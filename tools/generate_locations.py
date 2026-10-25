#!/usr/bin/env python3
"""
Generate Godot LocationData .tres from the `Locations` sheet of
tools/Machiavelli Roguelike.xlsx into data/locations/.

  Locations: Name | Type | Description

  Name         the location's name, shown on its card. Required. Its id is the
               name slugified ("Coin Rush" -> coin_rush), and the id is how the
               game finds the rule that makes it work (MatchLocation.ID_NAMES).
  Type         "On reveal" or "Ongoing". Required.
  Description  the rule text on the card. Required. A full stop is added if
               missing, since the game prints it inside sentences.

The sheet says what a location is CALLED and what its card SAYS; what it DOES is
code. A row whose id has no rule in code is still written, and the game leaves
it out of the pool with a warning — so a new location can be authored before
it is built. Deleting a row takes that location out of the game.

The output folder is REWRITTEN, not added to: a row deleted from the sheet takes
its .tres with it. Its README.md is left alone.

The workbook is only ever READ here. Never save it back with openpyxl: a
round-trip drops whatever openpyxl does not understand (charts, table styles).

  python3 tools/generate_locations.py           # write data/locations/*.tres
  python3 tools/generate_locations.py --list    # print the parse, write nothing
  python3 tools/generate_locations.py --check   # exit 1 if data/ is out of date
"""

import argparse
import glob
import os
import re
import sys

import openpyxl

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX_PATH = os.path.join(PROJECT_ROOT, "tools", "Machiavelli Roguelike.xlsx")
SHEET = "Locations"
OUT_DIR = os.path.join(PROJECT_ROOT, "data", "locations")
SCRIPT_RES = "res://scripts/resources/location_data.gd"

KINDS = {"on reveal": "on_reveal", "ongoing": "ongoing"}
REQUIRED = ("Name", "Type", "Description")


def slugify(name: str) -> str:
    s = str(name).strip().lower().replace("'", "")
    s = re.sub(r"[^a-z0-9]+", "_", s)
    return s.strip("_")


def gd_str(s: str) -> str:
    return s.replace("\\", "\\\\").replace('"', '\\"').replace("\n", " ").replace("\r", " ")


def clean(v) -> str:
    return re.sub(r"\s+", " ", "" if v is None else str(v)).strip()


def read_rows(path: str = XLSX_PATH) -> list:
    """Every location in the sheet, in sheet order. Raises ValueError on a bad row."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if SHEET not in wb.sheetnames:
        raise ValueError("no '%s' sheet in %s (has: %s)" % (SHEET, path, ", ".join(wb.sheetnames)))
    rows = wb[SHEET].iter_rows(values_only=True)
    header = [clean(h) for h in next(rows)]
    missing = [h for h in REQUIRED if h not in header]
    if missing:
        raise ValueError("'%s' sheet is missing column(s): %s" % (SHEET, ", ".join(missing)))
    col = {h: header.index(h) for h in REQUIRED}

    out, errors, seen = [], [], {}
    for n, row in enumerate(rows, start=2):
        cell = {h: clean(row[col[h]]) if col[h] < len(row) else "" for h in REQUIRED}
        if not any(cell.values()):
            continue  # blank row
        where = "%s row %d" % (SHEET, n)
        if not cell["Name"]:
            errors.append("%s: no Name" % where)
            continue
        loc_id = slugify(cell["Name"])
        if loc_id in seen:
            errors.append("%s: '%s' has the same id (%s) as row %d" % (where, cell["Name"], loc_id, seen[loc_id]))
            continue
        seen[loc_id] = n
        kind = KINDS.get(cell["Type"].lower())
        if kind is None:
            errors.append("%s (%s): Type '%s' is not one of: On reveal, Ongoing" % (where, cell["Name"], cell["Type"]))
            continue
        text = cell["Description"]
        if not text:
            errors.append("%s (%s): no Description" % (where, cell["Name"]))
            continue
        if text[-1] not in ".!?":
            text += "."
        out.append({"id": loc_id, "title": cell["Name"], "kind": kind, "text": text, "order": len(out)})
    if errors:
        raise ValueError("\n".join(errors))
    return out


def render(loc: dict) -> str:
    return (
        '[gd_resource type="Resource" script_class="LocationData" load_steps=2 format=3]\n'
        "\n"
        '[ext_resource type="Script" path="%s" id="1_loc"]\n'
        "\n"
        "[resource]\n"
        'script = ExtResource("1_loc")\n'
        'id = &"%s"\n'
        'title = "%s"\n'
        'kind = "%s"\n'
        'text = "%s"\n'
        "order = %d\n"
    ) % (SCRIPT_RES, loc["id"], gd_str(loc["title"]), loc["kind"], gd_str(loc["text"]), loc["order"])


def wanted_files(locs: list) -> dict:
    return {os.path.join(OUT_DIR, loc["id"] + ".tres"): render(loc) for loc in locs}


def existing_files() -> dict:
    out = {}
    for path in sorted(glob.glob(os.path.join(OUT_DIR, "*.tres"))):
        with open(path, encoding="utf-8") as f:
            out[path] = f.read()
    return out


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--list", action="store_true", help="print the parse, write nothing")
    ap.add_argument("--check", action="store_true", help="exit 1 if data/locations is out of date")
    args = ap.parse_args()

    try:
        locs = read_rows()
    except ValueError as e:
        print("generate_locations: %s" % e, file=sys.stderr)
        return 1

    if args.list:
        for loc in locs:
            print("%-12s %-10s %-12s %s" % (loc["id"], loc["kind"], loc["title"], loc["text"]))
        return 0

    want, have = wanted_files(locs), existing_files()
    rel = lambda p: os.path.relpath(p, PROJECT_ROOT)  # noqa: E731
    if args.check:
        stale = sorted(rel(p) for p in set(want) | set(have) if want.get(p) != have.get(p))
        if stale:
            print("data/locations is out of date with the sheet; run python3 tools/generate_locations.py")
            for p in stale:
                print("  " + p)
            return 1
        print("data/locations matches the sheet (%d locations)" % len(want))
        return 0

    os.makedirs(OUT_DIR, exist_ok=True)
    for path in set(have) - set(want):
        os.remove(path)
        print("removed %s" % rel(path))
    for path, body in want.items():
        if have.get(path) != body:
            with open(path, "w", encoding="utf-8", newline="\n") as f:
                f.write(body)
            print("wrote %s" % rel(path))
    print("%d locations" % len(want))
    return 0


if __name__ == "__main__":
    sys.exit(main())

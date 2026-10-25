#!/usr/bin/env python3
"""
Generate Godot InstanceData .tres from the `Instances` sheet of
tools/Machiavelli Roguelike.xlsx into data/instances/.

  Instances: Name | Type | Description | Effect

  Name         the instance's name, shown on its card. Required. Its id is the
               name slugified ("Coin Rush" -> coin_rush).
  Type         "Reveal" or "Ongoing" ("On reveal" still reads as Reveal). Required.
  Description  the rule text on the card. Required. A full stop is added if
               missing, since the game prints it inside sentences.
  Effect       what the instance DOES: commands separated by `;`, e.g.
               `discard choose 2; draw 2`. The commands are listed in
               data/instances/README.md. Optional: a row with no Effect is still
               written, and the game leaves it out of the pool with a warning —
               so a new instance can be named before its effect is written.

Every Effect is checked here against the command list (an unknown command is an
error, so a typo never reaches the game unnoticed); the game parses it in full
(InstanceEffect) and the smoke test plays every instance. Deleting a row takes
that instance out of the game.

The output folder is REWRITTEN, not added to: a row deleted from the sheet takes
its .tres with it. Its README.md is left alone.

The workbook is only ever READ here. Never save it back with openpyxl: a
round-trip drops whatever openpyxl does not understand (charts, table styles).

  python3 tools/generate_instances.py           # write data/instances/*.tres
  python3 tools/generate_instances.py --list    # print the parse, write nothing
  python3 tools/generate_instances.py --check   # exit 1 if data/ is out of date
"""

import argparse
import glob
import os
import re
import sys

import openpyxl

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX_PATH = os.path.join(PROJECT_ROOT, "tools", "Machiavelli Roguelike.xlsx")
SHEET = "Instances"
OUT_DIR = os.path.join(PROJECT_ROOT, "data", "instances")
SCRIPT_RES = "res://scripts/resources/instance_data.gd"

KINDS = {"reveal": "reveal", "on reveal": "reveal", "ongoing": "ongoing"}
REQUIRED = ("Name", "Type", "Description")
OPTIONAL = ("Effect",)
# The commands an Effect may use (InstanceEffect.VERB_NAMES in the game).
COMMANDS = ("draw", "draw_joker", "discard", "give", "jumble", "cost", "energy",
            "max_energy", "draw_bonus", "coins_per_card", "open_decks", "destroy")


def slugify(name: str) -> str:
    s = str(name).strip().lower().replace("'", "")
    s = re.sub(r"[^a-z0-9]+", "_", s)
    return s.strip("_")


def gd_str(s: str) -> str:
    return s.replace("\\", "\\\\").replace('"', '\\"').replace("\n", " ").replace("\r", " ")


def clean(v) -> str:
    return re.sub(r"\s+", " ", "" if v is None else str(v)).strip()


def read_rows(path: str = XLSX_PATH) -> list:
    """Every instance in the sheet, in sheet order. Raises ValueError on a bad row."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if SHEET not in wb.sheetnames:
        raise ValueError("no '%s' sheet in %s (has: %s)" % (SHEET, path, ", ".join(wb.sheetnames)))
    rows = wb[SHEET].iter_rows(values_only=True)
    header = [clean(h) for h in next(rows)]
    missing = [h for h in REQUIRED if h not in header]
    if missing:
        raise ValueError("'%s' sheet is missing column(s): %s" % (SHEET, ", ".join(missing)))
    col = {h: header.index(h) for h in REQUIRED + OPTIONAL if h in header}

    out, errors, seen = [], [], {}
    for n, row in enumerate(rows, start=2):
        cell = {h: clean(row[col[h]]) if h in col and col[h] < len(row) else ""
                for h in REQUIRED + OPTIONAL}
        if not any(cell.values()):
            continue  # blank row
        where = "%s row %d" % (SHEET, n)
        if not cell["Name"]:
            errors.append("%s: no Name" % where)
            continue
        loc_id = slugify(cell["Name"])
        if loc_id in seen:
            errors.append("%s: '%s' has the same id (%s) as row %d" % (where, cell["Name"], loc_id, seen[loc_id]))
            continue
        seen[loc_id] = n
        kind = KINDS.get(cell["Type"].lower())
        if kind is None:
            errors.append("%s (%s): Type '%s' is not one of: Reveal, Ongoing" % (where, cell["Name"], cell["Type"]))
            continue
        text = cell["Description"]
        if not text:
            errors.append("%s (%s): no Description" % (where, cell["Name"]))
            continue
        if text[-1] not in ".!?":
            text += "."
        effect = cell["Effect"]
        for command in (c.strip() for c in effect.split(";")):
            if command and command.split()[0].lower() not in COMMANDS:
                errors.append("%s (%s): Effect command '%s' is not one of: %s"
                              % (where, cell["Name"], command, ", ".join(COMMANDS)))
        out.append({"id": loc_id, "title": cell["Name"], "kind": kind, "text": text,
                    "effect": effect, "order": len(out)})
    if errors:
        raise ValueError("\n".join(errors))
    return out


def render(loc: dict) -> str:
    return (
        '[gd_resource type="Resource" script_class="InstanceData" load_steps=2 format=3]\n'
        "\n"
        '[ext_resource type="Script" path="%s" id="1_loc"]\n'
        "\n"
        "[resource]\n"
        'script = ExtResource("1_loc")\n'
        'id = &"%s"\n'
        'title = "%s"\n'
        'kind = "%s"\n'
        'text = "%s"\n'
        'effect = "%s"\n'
        "order = %d\n"
    ) % (SCRIPT_RES, loc["id"], gd_str(loc["title"]), loc["kind"], gd_str(loc["text"]),
         gd_str(loc["effect"]), loc["order"])


def wanted_files(locs: list) -> dict:
    return {os.path.join(OUT_DIR, loc["id"] + ".tres"): render(loc) for loc in locs}


def existing_files() -> dict:
    out = {}
    for path in sorted(glob.glob(os.path.join(OUT_DIR, "*.tres"))):
        with open(path, encoding="utf-8") as f:
            out[path] = f.read()
    return out


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--list", action="store_true", help="print the parse, write nothing")
    ap.add_argument("--check", action="store_true", help="exit 1 if data/instances is out of date")
    args = ap.parse_args()

    try:
        locs = read_rows()
    except ValueError as e:
        print("generate_instances: %s" % e, file=sys.stderr)
        return 1

    if args.list:
        for loc in locs:
            print("%-12s %-8s %-12s %-40s %s" % (loc["id"], loc["kind"], loc["title"],
                                                 loc["effect"] or "(no effect yet)", loc["text"]))
        return 0

    want, have = wanted_files(locs), existing_files()
    rel = lambda p: os.path.relpath(p, PROJECT_ROOT)  # noqa: E731
    if args.check:
        stale = sorted(rel(p) for p in set(want) | set(have) if want.get(p) != have.get(p))
        if stale:
            print("data/instances is out of date with the sheet; run python3 tools/generate_instances.py")
            for p in stale:
                print("  " + p)
            return 1
        print("data/instances matches the sheet (%d instances)" % len(want))
        return 0

    os.makedirs(OUT_DIR, exist_ok=True)
    for path in set(have) - set(want):
        os.remove(path)
        print("removed %s" % rel(path))
    for path, body in want.items():
        if have.get(path) != body:
            with open(path, "w", encoding="utf-8", newline="\n") as f:
                f.write(body)
            print("wrote %s" % rel(path))
    print("%d instances" % len(want))
    return 0


if __name__ == "__main__":
    sys.exit(main())
